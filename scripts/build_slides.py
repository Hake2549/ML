"""สร้างไฟล์นำเสนอ .pptx 17 หน้า ตามแผนใน docs/Slide-Plan.md

รูปกราฟดึงมาจาก outputs/figures/ ซึ่งเป็นผลการรันจริง
วิธีรัน (จาก root ของ repo):  python scripts/build_slides.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / 'outputs' / 'figures'
OUT = ROOT / 'slides' / 'Hotel-Booking-Cancellation-Prediction.pptx'
OUT.parent.mkdir(exist_ok=True)

NAVY = RGBColor(0x0F, 0x20, 0x38)
GOLD = RGBColor(0xE0, 0xA4, 0x58)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2C, 0x3D)
MUTED = RGBColor(0x6B, 0x7A, 0x8C)
CARD = RGBColor(0xF2, 0xF5, 0xF8)
RED = RGBColor(0xB3, 0x3A, 0x3A)
GREEN = RGBColor(0x1E, 0x7A, 0x5E)
FONT = 'Tahoma'

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)                      # ขอบซ้าย/ขวา
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def add_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY if dark else WHITE
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space_after=6, line_spacing=1.15, anchor=MSO_ANCHOR.TOP):
    """runs = str | list[str] | list[list[(text, {opts})]] (หนึ่งรายการ = หนึ่งย่อหน้า)"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        parts = [(para, {})] if isinstance(para, str) else para
        for t, opt in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get('size', size))
            f.bold = opt.get('bold', bold)
            f.color.rgb = opt.get('color', color)
    return box


def card(slide, x, y, w, h, fill=CARD, line=None, radius=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if radius and len(shp.adjustments):
        shp.adjustments[0] = 0.06
    return shp


def header(slide, number, title, kicker=None):
    """เลขหน้าในวงกลมทอง + หัวเรื่อง (motif ที่ใช้ซ้ำทุกหน้า)"""
    d = Inches(0.62)
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, M, Inches(0.42), d, d)
    circ.fill.solid()
    circ.fill.fore_color.rgb = GOLD
    circ.line.fill.background()
    circ.shadow.inherit = False
    tf = circ.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(20), True, NAVY
    tx = M + d + Inches(0.25)
    if kicker:
        text(slide, tx, Inches(0.3), W - tx - M, Inches(0.3), kicker, size=11, color=GOLD, bold=True)
        text(slide, tx, Inches(0.68), W - tx - M, Inches(0.5), title, size=26, color=NAVY, bold=True,
             space_after=0, line_spacing=1.0)
    else:
        text(slide, tx, Inches(0.5), W - tx - M, Inches(0.55), title, size=26, color=NAVY, bold=True,
             space_after=0, line_spacing=1.0)


def bullets(slide, x, y, w, items, size=14, gap=10):
    paras = []
    for it in items:
        parts = [(it, {})] if isinstance(it, str) else it
        paras.append([('•  ', {'color': GOLD, 'bold': True})] + list(parts))
    return text(slide, x, y, w, Inches(0.4), paras, size=size, space_after=gap, line_spacing=1.25)


def picture(slide, name, x, y, max_w, max_h):
    path = FIG / name
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(str(path), int(x + (max_w - w) / 2), int(y + (max_h - h) / 2),
                                    width=w, height=h)


def stat(slide, x, y, w, h, value, label, value_color=NAVY, fill=CARD, value_size=26):
    card(slide, x, y, w, h, fill=fill)
    text(slide, x + Inches(0.12), y + Inches(0.12), w - Inches(0.24), h - Inches(0.62), value,
         size=value_size, color=value_color, bold=True, align=PP_ALIGN.CENTER, space_after=0,
         line_spacing=1.0)
    text(slide, x + Inches(0.12), y + h - Inches(0.38), w - Inches(0.24), Inches(0.3), label,
         size=11, color=MUTED, align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)


def table(slide, x, y, w, rows, col_w, head_fill=NAVY, row_h=Inches(0.42), size=13,
          highlight=None):
    """rows[0] = หัวตาราง · highlight = index ของแถวที่ต้องการเน้น"""
    total = sum(col_w)
    col_w = [int(c / total * w) for c in col_w]
    for ri, row in enumerate(rows):
        cy = y + row_h * ri
        if ri == 0:
            card(slide, x, cy, w, row_h, fill=head_fill, radius=False)
        elif highlight is not None and ri == highlight:
            card(slide, x, cy, w, row_h, fill=RGBColor(0xFD, 0xF3, 0xE3), radius=False)
        elif ri % 2 == 0:
            card(slide, x, cy, w, row_h, fill=CARD, radius=False)
        cx = x
        for ci, cell in enumerate(row):
            val = cell if isinstance(cell, str) else cell[0]
            bold = ri == 0 or (not isinstance(cell, str) and cell[1])
            color = WHITE if ri == 0 else (NAVY if bold else INK)
            align = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            text(slide, cx + Inches(0.12), cy + Inches(0.08), col_w[ci] - Inches(0.24),
                 row_h - Inches(0.1), val, size=size, color=color, bold=bold, align=align,
                 space_after=0, line_spacing=1.0)
            cx += col_w[ci]
    return y + row_h * len(rows)


def note(slide, txt, y=None, color=MUTED, size=12):
    y = y if y is not None else H - Inches(0.75)
    return text(slide, M, y, W - 2 * M, Inches(0.45), txt, size=size, color=color, line_spacing=1.2)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


# ---------------------------------------------------------------- 1 title
s = add_slide(dark=True)
text(s, M, Inches(1.5), W - 2 * M, Inches(0.4), 'โครงงานรายวิชา Machine Learning',
     size=15, color=GOLD, bold=True)
text(s, M, Inches(2.0), Inches(9.6), Inches(1.6),
     'Hotel Booking\nCancellation Prediction', size=44, color=WHITE, bold=True, line_spacing=1.05)
text(s, M, Inches(3.75), Inches(9.6), Inches(0.5),
     'ทำนายความเสี่ยงที่การจองโรงแรมจะถูกยกเลิก ด้วยข้อมูลการจองจริง 119,208 รายการ',
     size=16, color=RGBColor(0xC7, 0xD2, 0xDE))
names_l = ['นายชัยพร ศรีพราย  66114540168', 'นางสาวกวิณชฎา มีศรี  67114540015',
           'นายธนากร ไชยรัตน์  67114540244']
names_r = ['นายรัชชานนท์ รื่นจิตต์  67114540446', 'นายราชภัฏ บุญรอด  67114540455']
text(s, M, Inches(4.9), Inches(4.6), Inches(1.4), names_l, size=13,
     color=RGBColor(0xE4, 0xEA, 0xF1), space_after=8)
text(s, Inches(5.5), Inches(4.9), Inches(4.6), Inches(1.4), names_r, size=13,
     color=RGBColor(0xE4, 0xEA, 0xF1), space_after=8)
stat(s, Inches(10.3), Inches(2.0), Inches(2.4), Inches(1.5), '18', 'features ที่ใช้เทรน',
     value_color=GOLD, fill=RGBColor(0x18, 0x2E, 0x4B))
stat(s, Inches(10.3), Inches(3.7), Inches(2.4), Inches(1.5), '0.847', 'F1-score บนชุดทดสอบ',
     value_color=GOLD, fill=RGBColor(0x18, 0x2E, 0x4B))
notes(s, 'สวัสดีครับ/ค่ะ กลุ่มเรานำเสนอโครงงานทำนายการยกเลิกการจองโรงแรม')

# ---------------------------------------------------------------- 2 why
s = add_slide()
header(s, 2, 'ยกเลิก 37% คือรายได้ที่หายและพยากรณ์ห้องที่ผิดพลาด', kicker='PROBLEM')
bullets(s, M, Inches(1.65), Inches(7.4), [
    'ยกเลิกกระชั้นชิด → ขายห้องต่อไม่ทัน → รายได้หายทันที',
    'โรงแรมต้องทำ Overbooking ชดเชย ซึ่งถ้าประเมินผิดจะเสียทั้งเงินและชื่อเสียง',
    [('Input = ', {'bold': True, 'color': NAVY}), ('ข้อมูลการจอง ณ วันที่จอง  →  ', {}),
     ('Output = ', {'bold': True, 'color': NAVY}), ('ยกเลิก (1) / ไม่ยกเลิก (0)', {})],
    'ขอบเขต: โรงแรม 2 แห่งในโปรตุเกส ก.ค. 2015 – ส.ค. 2017',
], size=15)
card(s, Inches(8.5), Inches(1.65), Inches(4.2), Inches(2.5), fill=NAVY)
text(s, Inches(8.5), Inches(2.0), Inches(4.2), Inches(1.2), '37%', size=60, color=GOLD, bold=True,
     align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
text(s, Inches(8.7), Inches(3.25), Inches(3.8), Inches(0.7),
     'ของการจองทั้งหมดถูกยกเลิก\n(44,224 จาก 119,390 รายการ)', size=13, color=WHITE,
     align=PP_ALIGN.CENTER, line_spacing=1.2)
note(s, 'ปัญหานี้เป็น Binary Classification ที่มีมูลค่าทางธุรกิจวัดได้จริง', y=Inches(4.5))
notes(s, 'เน้นว่าเป็นปัญหาที่วัดมูลค่าทางธุรกิจได้ ไม่ใช่แค่โจทย์ในห้องเรียน')

# ---------------------------------------------------------------- 3 imbalance
s = add_slide()
header(s, 3, 'ทายว่า "ไม่ยกเลิก" ทุกรายการ ก็ได้ Accuracy 62.9% แล้ว', kicker='DATA UNDERSTANDING')
picture(s, '3_target_distribution.png', Inches(7.6), Inches(1.6), Inches(5.1), Inches(4.4))
bullets(s, M, Inches(1.75), Inches(6.6), [
    'ข้อมูลดิบ 119,390 แถว × 32 คอลัมน์',
    'ยกเลิก 44,224 รายการ : ไม่ยกเลิก 75,166 รายการ',
    'สัดส่วนคลาสไม่สมดุล 37 : 63',
], size=15)
card(s, M, Inches(3.5), Inches(6.6), Inches(1.5), fill=RGBColor(0xFD, 0xF3, 0xE3))
text(s, M + Inches(0.3), Inches(3.75), Inches(6.0), Inches(1.1),
     [[('ประโยคสำคัญ: ', {'bold': True, 'color': GOLD})],
      [('“นี่คือเหตุผลที่เราไม่ใช้ Accuracy เป็นตัวชี้วัดหลัก ', {'bold': True, 'color': NAVY}),
       ('เพราะโมเดลที่ไม่ได้เรียนรู้อะไรเลยก็ทำคะแนนนี้ได้”', {'bold': True, 'color': NAVY})]],
     size=14, line_spacing=1.25)
notes(s, 'ปูทางไปหน้า 11 เรื่องการเลือก metric')

# ---------------------------------------------------------------- 4 EDA
s = add_slide()
header(s, 4, 'ประเภทมัดจำและช่องทางการจองต่างกันอย่างชัดเจน', kicker='EDA')
picture(s, '3_eda_cancel_rate_by_category.png', Inches(4.3), Inches(1.5), Inches(8.4), Inches(5.2))
bullets(s, M, Inches(1.75), Inches(3.4), [
    [('Non Refund ', {'bold': True, 'color': NAVY}), ('ยกเลิกเกือบ 100% ทั้งที่จ่ายเงินไปแล้ว', {})],
    [('Groups ', {'bold': True, 'color': NAVY}), ('ยกเลิกสูงกว่าการจองแบบอื่นชัดเจน', {})],
    'เส้นประแดง = อัตรายกเลิกเฉลี่ยทั้งชุดข้อมูล',
], size=14)
note(s, 'ทุกกราฟคำนวณจากข้อมูลจริงทั้ง 119,390 แถว', y=Inches(6.75))
notes(s, 'พูดแค่ 2 จุดที่สะดุดตาที่สุด ไม่ต้องไล่ทุกกราฟ')

# ---------------------------------------------------------------- 5 cleaning
s = add_slide()
header(s, 5, 'ทำความสะอาดข้อมูลโดยเสียข้อมูลเพียง 0.15%', kicker='DATA PREPARATION')
end_y = table(s, M, Inches(1.7), Inches(8.0), [
    ['ปัญหาที่พบ', 'วิธีจัดการ', 'แถวที่เสีย'],
    ['children ว่าง 4 ค่า', 'เติมด้วย 0', '0'],
    ['country ว่าง 488 ค่า', 'แทนด้วย Unknown', '0'],
    ['agent / company ว่าง', 'แปลงเป็น flag มี/ไม่มี', '0'],
    ['ไม่มีผู้เข้าพัก', 'ลบแถวทิ้ง', ('180', True)],
    ['adr ติดลบ / เกิน 1,000', 'ลบแถวทิ้ง', ('2', True)],
], col_w=[3.4, 3.2, 1.4])
stat(s, Inches(9.0), Inches(1.7), Inches(3.7), Inches(1.3), '119,390 → 119,208',
     'จำนวนแถวก่อน → หลังคลีน', value_color=NAVY, value_size=19)
stat(s, Inches(9.0), Inches(3.2), Inches(3.7), Inches(1.3), '182 แถว', 'ลบทิ้งทั้งหมด คิดเป็น 0.15%', value_color=GOLD)
text(s, Inches(9.0), Inches(4.75), Inches(3.7), Inches(1.2),
     'อัตรายกเลิกหลังคลีน 37.08%\nแทบไม่ต่างจาก 37.04% ก่อนคลีน\n→ การลบแถวไม่ทำให้สัดส่วนคลาสบิดเบือน',
     size=12, color=MUTED, line_spacing=1.3)
text(s, M, end_y + Inches(0.3), Inches(8.0), Inches(0.6),
     'ค่าว่างใน agent / company ตามเปเปอร์ต้นฉบับแปลว่า “ไม่เกี่ยวข้อง” ไม่ใช่ข้อมูลหาย '
     'จึงทำเป็น flag แทนการเติมค่า', size=12, color=MUTED, line_spacing=1.25)
notes(s, 'ย้ำว่าเลือกวิธีจัดการจากความหมายของข้อมูล ไม่ได้เติมค่ามั่ว')

# ---------------------------------------------------------------- 6 leakage
s = add_slide()
header(s, 6, 'ตัด 3 คอลัมน์ที่แอบรู้คำตอบ ก่อนแบ่งข้อมูล', kicker='DATA LEAKAGE')
card(s, M, Inches(1.65), Inches(6.1), Inches(2.5), fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, M + Inches(0.3), Inches(1.85), Inches(5.5), Inches(0.3), 'คอลัมน์ที่ตัดออก', size=13,
     color=RED, bold=True)
bullets(s, M + Inches(0.3), Inches(2.25), Inches(5.5), [
    'reservation_status — บอกผลลัพธ์ตรงๆ',
    'reservation_status_date — วันที่ผลลัพธ์เกิดขึ้น',
    'assigned_room_type — ห้องถูกกำหนดตอนเช็คอิน',
], size=13, gap=7)
text(s, M + Inches(0.3), Inches(3.55), Inches(5.5), Inches(0.45),
     [[('หลักฐาน: ', {'bold': True, 'color': RED}),
       ('การจองที่ถูกเปลี่ยนห้องมีอัตรายกเลิก 0% → รู้ผลหลังเช็คอินแล้ว', {'bold': True})]],
     size=13, color=NAVY, line_spacing=1.2)
text(s, Inches(7.1), Inches(1.65), Inches(5.6), Inches(0.3), 'แบ่งข้อมูล 2 แบบ', size=13,
     color=NAVY, bold=True)
table(s, Inches(7.1), Inches(2.05), Inches(5.6), [
    ['วิธีแบ่ง', 'Train', 'Test'],
    ['Stratified Random 80/20', '95,366', '23,842'],
    ['Time-based (ตามเวลา)', '95,190', '24,018'],
], col_w=[3.0, 1.3, 1.3])
text(s, Inches(7.1), Inches(3.5), Inches(5.6), Inches(0.9),
     'Time split: train มีอัตรายกเลิก 36.09% แต่ test มี 40.99%\n→ พฤติกรรมเปลี่ยนตามเวลาจริง (ดูหน้า 14)',
     size=12, color=MUTED, line_spacing=1.3)
note(s, 'encoding · scaling · feature selection ทั้งหมดอยู่ใน Pipeline จึงเรียนรู้จาก train ของแต่ละ fold เท่านั้น — Test set ถูกเปิดใช้ครั้งเดียวตอนท้าย',
     y=Inches(4.55))
notes(s, 'หน้านี้คืออันที่อาจารย์ชอบถาม ให้เน้นหลักฐาน 0%')

# ---------------------------------------------------------------- 7 feature engineering
s = add_slide()
header(s, 7, 'สร้าง 9 features ใหม่ — ทุกตัวมีสมมติฐานรองรับ', kicker='FEATURE ENGINEERING')
picture(s, '5_engineered_features_cancel_rate.png', M, Inches(1.55), Inches(7.5), Inches(5.2))
text(s, Inches(8.4), Inches(1.6), Inches(4.3), Inches(0.3),
     'ช่วงห่างของอัตรายกเลิกระหว่างกลุ่ม', size=13, color=NAVY, bold=True)
table(s, Inches(8.4), Inches(2.0), Inches(4.3), [
    ['Feature ใหม่', 'ช่วงห่าง'],
    ['cancel_history_ratio', ('81.6%', True)],
    ['lead_time_group', ('58.1%', True)],
    ['total_nights', '52.8%'],
    ['is_domestic', '33.1%'],
], col_w=[3.0, 1.3], highlight=1)
text(s, Inches(8.4), Inches(4.35), Inches(4.3), Inches(1.6),
     'ยิ่งช่วงห่างกว้าง feature ยิ่งแยกกลุ่มคนที่จะยกเลิกออกจากคนที่ไม่ยกเลิกได้ดี\n\n'
     'ตั้งสมมติฐานก่อน แล้วค่อยพิสูจน์ด้วยกราฟ ไม่ได้สร้างมั่วแล้วลองดู',
     size=12, color=MUTED, line_spacing=1.3)
notes(s, 'เน้นว่าทุก feature มีเหตุผลรองรับก่อนสร้าง')

# ---------------------------------------------------------------- 8 selection
s = add_slide()
header(s, 8, '37 features → 18 features ด้วยการโหวต 2 ใน 3', kicker='FEATURE SELECTION')
labels = [('Mutual Information', 'Filter — สถิติล้วน', '5_mutual_information.png'),
          ('RF Importance', 'Embedded — จากตัวโมเดล', '5_random_forest_importance.png'),
          ('Permutation Importance', 'Model-agnostic — วัดบน validation', '5_permutation_importance.png')]
x = M
cw = Inches(3.9)
for title_, sub, fname in labels:
    text(s, x, Inches(1.5), cw, Inches(0.25), title_, size=13, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER, space_after=0)
    text(s, x, Inches(1.78), cw, Inches(0.22), sub, size=11, color=MUTED, align=PP_ALIGN.CENTER,
         space_after=0)
    picture(s, fname, x, Inches(2.1), cw, Inches(3.3))
    x += cw + Inches(0.32)
stat(s, M, Inches(5.65), Inches(2.6), Inches(1.15), '14 ตัว', 'ได้ครบ 3 โหวต')
stat(s, M + Inches(2.85), Inches(5.65), Inches(2.6), Inches(1.15), '4 ตัว', 'ได้ 2 โหวต')
stat(s, M + Inches(5.7), Inches(5.65), Inches(2.6), Inches(1.15), '18 ตัว', 'เก็บไว้ใช้เทรน',
     value_color=GOLD, fill=RGBColor(0xFD, 0xF3, 0xE3))
text(s, M + Inches(8.65), Inches(5.7), Inches(3.75), Inches(1.05),
     'RF Importance มีอคติเข้าข้างตัวแปรที่มีค่าหลากหลาย จึงใช้ Permutation Importance '
     'ที่วัดบน validation set เป็นตัวถ่วง', size=12, color=MUTED, line_spacing=1.25)
notes(s, 'ถ้าโดนถามเรื่อง overfitting ให้เปิดสไลด์สำรองหน้าสุดท้าย (Spearman 0.985)')

# ---------------------------------------------------------------- 9 A/B/C
s = add_slide()
header(s, 9, 'ใช้ feature น้อยลงครึ่งหนึ่ง แต่ผลเท่าเดิมและเทรนเร็วขึ้น', kicker='ก่อน – หลังคัดเลือก')
table(s, M, Inches(1.7), Inches(7.4), [
    ['ชุด Feature', 'จำนวน', 'F1 (CV)', 'ROC-AUC'],
    ['A: ตัวแปรเดิม', '28', '0.8103 ± 0.0049', '0.9420'],
    ['B: เดิม + ที่สร้างใหม่', '37', '0.8127 ± 0.0042', '0.9428'],
    [('C: คัดเลือกแล้ว', True), ('18', True), ('0.8112 ± 0.0027', True), ('0.9417', True)],
], col_w=[3.0, 1.1, 2.2, 1.4], row_h=Inches(0.46), highlight=3)
card(s, M, Inches(3.85), Inches(7.4), Inches(1.5), fill=RGBColor(0xFD, 0xF3, 0xE3))
text(s, M + Inches(0.3), Inches(4.05), Inches(6.8), Inches(1.1),
     [[('ผลต่างระหว่าง B กับ C = 0.0015 ', {'bold': True, 'color': NAVY}),
       ('ซึ่งน้อยกว่าค่า std (0.0042)', {})],
      [('→ ถือว่าไม่ต่างกันอย่างมีนัยสำคัญ แต่ C มีความแปรปรวนต่ำที่สุดและเทรนเร็วที่สุด จึงเลือก C',
        {'bold': True, 'color': NAVY})]], size=14, line_spacing=1.3)
picture(s, '5_performance_vs_k.png', Inches(8.3), Inches(1.7), Inches(4.4), Inches(3.4))
text(s, Inches(8.3), Inches(5.25), Inches(4.4), Inches(0.9),
     'กราฟตอบคำถาม “ทำไมเลือก K = 18” — เพิ่ม feature เกินจุดนี้แล้วคะแนนไม่ขึ้นอีก',
     size=12, color=MUTED, align=PP_ALIGN.CENTER, line_spacing=1.25)
notes(s, 'ถ้าโดนถามว่าทำไม C ไม่ดีกว่า B ให้ตอบเรื่อง std')

# ---------------------------------------------------------------- 10 models
s = add_slide()
header(s, 10, 'Random Forest นำตั้งแต่ก่อน Tuning (F1 = 0.842)', kicker='MODEL DEVELOPMENT')
picture(s, '6_default_models_f1.png', Inches(6.9), Inches(1.6), Inches(5.8), Inches(3.6))
table(s, M, Inches(1.7), Inches(6.0), [
    ['โมเดล (ค่า default)', 'F1 (CV)', 'ช่องว่าง Train–CV'],
    ['Dummy (baseline)', '0.000', '0.000'],
    ['Logistic Regression', '0.763', '0.001'],
    ['Decision Tree', '0.791', ('0.203', True)],
    [('Random Forest', True), ('0.842', True), '0.152'],
    ['HistGradientBoosting', '0.811', '0.011'],
], col_w=[2.8, 1.4, 1.9], highlight=4)
card(s, M, Inches(4.7), Inches(6.0), Inches(1.3), fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, M + Inches(0.25), Inches(4.9), Inches(5.5), Inches(0.9),
     [[('Decision Tree overfit ชัดเจน ', {'bold': True, 'color': RED}),
       ('— ช่องว่าง train–CV ถึง 0.203', {'bold': True, 'color': NAVY})],
      [('Random Forest แก้ปัญหานี้ด้วยการรวมหลายต้นเข้าด้วยกัน', {})]],
     size=13, line_spacing=1.3)
note(s, 'ทุกโมเดลใช้ Stratified 5-Fold Cross-Validation บน train set เดียวกัน และใช้ Pipeline ชุดเดียวกัน', y=Inches(6.3))
notes(s, 'ชี้ให้เห็นว่า Dummy ได้ F1 = 0 ทั้งที่ Accuracy สูง')

# ---------------------------------------------------------------- 11 metrics
s = add_slide()
header(s, 11, 'Accuracy หลอกตา — Dummy ได้ 62.9% แต่ F1 = 0', kicker='METRICS')
card(s, M, Inches(1.65), Inches(6.0), Inches(1.5), fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, M + Inches(0.25), Inches(1.85), Inches(5.5), Inches(1.2),
     [[('False Negative', {'bold': True, 'color': RED, 'size': 15})],
      [('ทำนายว่าไม่ยกเลิก → แต่ยกเลิกจริง', {'size': 13})],
      [('ห้องว่าง ขายต่อไม่ทัน = รายได้หาย', {'size': 13, 'bold': True, 'color': NAVY})]],
     line_spacing=1.25, space_after=3)
card(s, Inches(6.9), Inches(1.65), Inches(5.8), Inches(1.5), fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, Inches(7.15), Inches(1.85), Inches(5.3), Inches(1.2),
     [[('False Positive', {'bold': True, 'color': RED, 'size': 15})],
      [('ทำนายว่ายกเลิก → แต่ลูกค้ามาจริง', {'size': 13})],
      [('ถ้าเผื่อ Overbooking ไว้ = ห้องไม่พอ เสียชื่อเสียง', {'size': 13, 'bold': True, 'color': NAVY})]],
     line_spacing=1.25, space_after=3)
text(s, M, Inches(3.45), Inches(12.1), Inches(0.35),
     'ทั้งสองแบบมีต้นทุนสูงพอกัน จึงเลือก F1-score เป็นตัวชี้วัดหลัก', size=16, color=NAVY, bold=True)
xs = M
for val, lab, col in [('F1-score', 'ตัวชี้วัดหลัก — สมดุล Precision กับ Recall', GOLD),
                      ('ROC-AUC', 'ความสามารถแยกคลาสโดยไม่ขึ้นกับ threshold', NAVY),
                      ('PR-AUC', 'เหมาะกับข้อมูลไม่สมดุลมากกว่า ROC-AUC', NAVY),
                      ('Confusion\nMatrix', 'เห็นจำนวน error แต่ละแบบจริง', NAVY)]:
    stat(s, xs, Inches(4.0), Inches(2.9), Inches(1.75), val, lab, value_color=col,
         fill=RGBColor(0xFD, 0xF3, 0xE3) if col == GOLD else CARD)
    xs += Inches(3.05)
note(s, 'Baseline: Dummy Classifier ทายคลาสที่พบบ่อยที่สุดเสมอ ได้ Accuracy 62.9% แต่ F1 = 0 คือจับคนที่จะยกเลิกไม่ได้เลยแม้แต่รายเดียว',
     y=Inches(6.1))
notes(s, 'ผูก metric เข้ากับต้นทุนธุรกิจจริง ไม่ใช่เลือกเพราะเห็นคนอื่นใช้')

# ---------------------------------------------------------------- 12 tuning
s = add_slide()
header(s, 12, 'Tuning ช่วยลด Overfitting มากกว่าจะเพิ่มคะแนน', kicker='HYPERPARAMETER TUNING')
table(s, M, Inches(1.7), Inches(7.2), [
    ['โมเดล', 'F1 ก่อน', 'F1 หลัง', 'ช่องว่าง Train–CV'],
    ['Logistic Regression', '0.7629', ('0.7784', True), '0.001 → 0.002'],
    [('Random Forest', True), '0.8422', '0.8419', ('0.152 → 0.122', True)],
    ['HistGradientBoosting', '0.8112', ('0.8344', True), '0.011 → 0.062'],
], col_w=[2.6, 1.3, 1.3, 2.2], row_h=Inches(0.46), highlight=2)
card(s, M, Inches(3.65), Inches(7.2), Inches(1.45), fill=RGBColor(0xFD, 0xF3, 0xE3))
text(s, M + Inches(0.28), Inches(3.85), Inches(6.6), Inches(1.05),
     [[('F1 ของ Random Forest แทบไม่ขึ้นเลย ', {'bold': True, 'color': NAVY}),
       ('แต่ช่องว่าง train–CV แคบลงจาก 0.152 เหลือ 0.122', {})],
      [('→ โมเดลจำข้อมูลน้อยลงและมีเสถียรภาพมากขึ้น จึงถือว่า tuning ได้ผล', {'bold': True, 'color': NAVY})]],
     size=14, line_spacing=1.3)
text(s, M, Inches(5.3), Inches(7.2), Inches(0.9),
     [[('ค่าที่เลือกใช้: ', {'bold': True, 'color': NAVY})],
      [('n_estimators=300 · max_depth=30 · min_samples_leaf=2 · max_features=log2 · '
        'class_weight=balanced_subsample', {'size': 12, 'color': MUTED})]], size=13, line_spacing=1.25)
picture(s, '7_validation_curve_rf_max_depth.png', Inches(8.1), Inches(1.7), Inches(4.6), Inches(3.4))
text(s, Inches(8.1), Inches(5.25), Inches(4.6), Inches(1.0),
     'Validation Curve: เส้น train พุ่งขึ้นแต่ validation นิ่ง = สัญญาณ overfitting ใช้เลือกค่า max_depth',
     size=12, color=MUTED, align=PP_ALIGN.CENTER, line_spacing=1.25)
notes(s, 'อย่ากลัวที่จะบอกว่า F1 ไม่ขึ้น อธิบายด้วยช่องว่าง train-CV แทน')

# ---------------------------------------------------------------- 13 test
s = add_slide()
header(s, 13, 'ปรับ threshold เป็น 0.46 เพิ่ม Recall ได้ 2.6 จุด', kicker='ผลบนชุดทดสอบ')
picture(s, '7_confusion_matrix.png', M, Inches(1.6), Inches(4.6), Inches(4.3))
xs = Inches(5.6)
for val, lab, col in [('0.8469', 'F1-score', GOLD), ('0.8351', 'Precision', NAVY),
                      ('0.8590', 'Recall', NAVY)]:
    stat(s, xs, Inches(1.75), Inches(2.25), Inches(1.4), val, lab, value_color=col,
         fill=RGBColor(0xFD, 0xF3, 0xE3) if col == GOLD else CARD)
    xs += Inches(2.4)
xs = Inches(5.6)
for val, lab in [('0.9567', 'ROC-AUC'), ('0.9367', 'PR-AUC'), ('0.8849', 'Accuracy')]:
    stat(s, xs, Inches(3.3), Inches(2.25), Inches(1.4), val, lab)
    xs += Inches(2.4)
card(s, Inches(5.6), Inches(4.9), Inches(7.1), Inches(1.35), fill=NAVY)
text(s, Inches(5.85), Inches(5.1), Inches(6.6), Inches(1.0),
     [[('Test set 23,842 แถว ถูกใช้เพียงครั้งเดียวตอนท้าย', {'bold': True, 'color': WHITE})],
      [('การเลือก feature · เลือกโมเดล · tuning · หา threshold ทั้งหมดใช้ Cross-Validation บน train เท่านั้น',
        {'color': RGBColor(0xC7, 0xD2, 0xDE), 'size': 13})]], size=14, line_spacing=1.3)
notes(s, 'threshold 0.46 หามาจาก out-of-fold predictions บน train ไม่ได้แตะ test')

# ---------------------------------------------------------------- 14 time split
s = add_slide()
header(s, 14, 'ถ้าทดสอบแบบทำนายอนาคตจริง F1 ตกจาก 0.847 เหลือ 0.725', kicker='จุดค้นพบสำคัญ')
table(s, M, Inches(1.75), Inches(7.6), [
    ['วิธีแบ่งข้อมูล (threshold 0.46)', 'F1', 'Precision', 'Recall'],
    [('Random split', True), ('0.8469', True), '0.8351', '0.8590'],
    [('Time-based split', True), ('0.7249', True), '0.7996', '0.6630'],
], col_w=[3.4, 1.2, 1.4, 1.3], row_h=Inches(0.5), highlight=2)
card(s, M, Inches(3.5), Inches(7.6), Inches(1.6), fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, M + Inches(0.3), Inches(3.72), Inches(7.0), Inches(1.2),
     [[('สาเหตุ: ', {'bold': True, 'color': RED}),
       ('ชุด test ที่แบ่งตามเวลามีอัตรายกเลิก 40.99% ขณะที่ train มี 36.09%', {'bold': True, 'color': NAVY})],
      [('พฤติกรรมลูกค้าเปลี่ยนไปตามเวลาจริง (distribution shift)', {})]],
     size=14, line_spacing=1.3)
stat(s, Inches(8.7), Inches(1.75), Inches(4.0), Inches(1.9), '−0.122', 'F1 ที่หายไปเมื่อทดสอบตามเวลา',
     value_color=RED, fill=RGBColor(0xFB, 0xEE, 0xEE))
text(s, Inches(8.7), Inches(3.9), Inches(4.0), Inches(2.0),
     [[('“การสุ่มแบ่งข้อมูลประเมินโมเดลสูงเกินจริง เราจึงรายงานตัวเลขทั้งสองแบบ”',
        {'bold': True, 'color': NAVY, 'size': 14})],
      [('ในการใช้งานจริง โมเดลต้องทำนายอนาคตจากข้อมูลอดีต จึงควร retrain เป็นระยะ',
        {'color': MUTED, 'size': 12})]], line_spacing=1.3, space_after=10)
notes(s, 'หน้านี้คือหน้าที่ทำให้งานดูลึกกว่ากลุ่มอื่น')

# ---------------------------------------------------------------- 15 deploy
s = add_slide()
header(s, 15, 'กรอกข้อมูลการจอง → ได้ความเสี่ยงทันที', kicker='DEPLOYMENT')
text(s, M, Inches(1.6), Inches(12.1), Inches(0.35),
     'เว็บแอป Streamlit — ผู้ใช้กรอกรายละเอียดการจอง ระบบคำนวณ 18 features แล้วแสดงความน่าจะเป็นที่จะยกเลิก',
     size=14, color=MUTED)
cases = [('99.8%', 'จองล่วงหน้า 300 วัน\nNon Refund + Groups', RED, 'เสี่ยงสูง'),
         ('25.9%', 'จองล่วงหน้า 60 วัน\nOnline TA', GOLD, 'เสี่ยงปานกลาง'),
         ('1.0%', 'จองล่วงหน้า 5 วัน\nDirect + ลูกค้าเก่า', GREEN, 'เสี่ยงต่ำ')]
x = M
for val, desc, col, tag in cases:
    card(s, x, Inches(2.2), Inches(3.9), Inches(2.6), fill=CARD)
    text(s, x, Inches(2.45), Inches(3.9), Inches(0.9), val, size=40, color=col, bold=True,
         align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
    text(s, x, Inches(3.35), Inches(3.9), Inches(0.3), tag, size=13, color=col, bold=True,
         align=PP_ALIGN.CENTER, space_after=0)
    text(s, x + Inches(0.3), Inches(3.8), Inches(3.3), Inches(0.8), desc, size=12, color=MUTED,
         align=PP_ALIGN.CENTER, line_spacing=1.3)
    x += Inches(4.1)
card(s, M, Inches(5.1), Inches(12.1), Inches(1.35), fill=NAVY)
text(s, M + Inches(0.3), Inches(5.3), Inches(11.5), Inches(1.0),
     [[('ข้อจำกัดของระบบ', {'bold': True, 'color': GOLD})],
      [('โมเดลบนเว็บใช้ HistGradientBoosting ขนาด 0.8 MB (F1 0.8334) แทน Random Forest 356.8 MB '
        'เพราะข้อจำกัดด้านหน่วยความจำของบริการฟรี · ผลทำนายอิงข้อมูลโปรตุเกสปี 2015–2017',
        {'color': RGBColor(0xC7, 0xD2, 0xDE), 'size': 12})]], size=14, line_spacing=1.3)
notes(s, 'ถ้าเน็ตในห้องใช้ได้ ให้เปิดเว็บจริงสาธิต 1 เคส')

# ---------------------------------------------------------------- 16 conclusion
s = add_slide()
header(s, 16, 'สรุปผล ข้อจำกัด และงานต่อไป', kicker='CONCLUSION')
cols = [
    ('ตอบโจทย์หรือไม่', GREEN, [
        'ทำนายได้ F1 0.8469 บนชุดทดสอบ',
        'สูงกว่า baseline (F1 = 0) อย่างชัดเจน',
        'ระบุปัจจัยเสี่ยงหลักได้: มัดจำ · เอเจนต์ · ระยะจองล่วงหน้า',
    ]),
    ('ข้อจำกัด', RED, [
        'ข้อมูลจากโปรตุเกส ปี 2015–2017 เท่านั้น',
        'ทดสอบตามเวลาได้ F1 ~0.72 ต่ำกว่าผลสุ่มแบ่ง',
        'มีแถวซ้ำ 31,994 แถวที่เลือกไม่ลบ',
    ]),
    ('งานต่อไป', NAVY, [
        'ใช้ time-based validation ตอน tuning',
        'ตั้ง threshold ตามต้นทุนธุรกิจจริง',
        'เพิ่ม SHAP อธิบายผลเป็นรายรายการ',
    ]),
]
x = M
for title_, col, items in cols:
    card(s, x, Inches(1.65), Inches(3.9), Inches(3.0), fill=CARD)
    text(s, x + Inches(0.28), Inches(1.9), Inches(3.3), Inches(0.3), title_, size=15, color=col, bold=True)
    bullets(s, x + Inches(0.28), Inches(2.35), Inches(3.35), items, size=12, gap=8)
    x += Inches(4.1)
card(s, M, Inches(4.95), Inches(12.1), Inches(1.5), fill=NAVY)
text(s, M + Inches(0.3), Inches(5.15), Inches(11.5), Inches(1.1),
     [[('สิ่งที่ได้เรียนรู้', {'bold': True, 'color': GOLD})],
      [('ระวัง Data Leakage ตั้งแต่ก่อนแบ่งข้อมูล · เลือก metric ให้ตรงกับต้นทุนของปัญหา · '
        'ใช้ Pipeline ป้องกันข้อมูลรั่วระหว่าง Cross-Validation · '
        'ตัวเลขที่ดูดีเกินจริงมักมาจากวิธีประเมินที่ไม่ตรงกับการใช้งานจริง',
        {'color': RGBColor(0xC7, 0xD2, 0xDE), 'size': 13})]], size=14, line_spacing=1.35)
notes(s, 'ปิดด้วยสิ่งที่ได้เรียนรู้ ไม่ใช่แค่ตัวเลข')

# ---------------------------------------------------------------- 17 Q&A
s = add_slide(dark=True)
text(s, M, Inches(2.3), Inches(7.6), Inches(1.0), 'ขอบคุณครับ / ค่ะ', size=44, color=WHITE, bold=True,
     space_after=0, line_spacing=1.0)
text(s, M, Inches(3.4), Inches(7.6), Inches(0.5), 'Q & A', size=26, color=GOLD, bold=True)
text(s, M, Inches(4.3), Inches(7.6), Inches(1.4),
     [[('โค้ดทั้งหมด (Colab / GitHub)', {'bold': True, 'color': WHITE, 'size': 14})],
      [('github.com/Hake2549/ML', {'color': RGBColor(0xC7, 0xD2, 0xDE), 'size': 14})]],
     line_spacing=1.4)
card(s, Inches(8.4), Inches(2.1), Inches(4.3), Inches(3.4), fill=RGBColor(0x18, 0x2E, 0x4B))
text(s, Inches(8.7), Inches(2.35), Inches(3.7), Inches(0.3), 'สไลด์สำรอง (เปิดเมื่อถูกถาม)',
     size=13, color=GOLD, bold=True)
bullets(s, Inches(8.7), Inches(2.8), Inches(3.7), [
    [('RF Importance ก่อน–หลัง tune', {'color': WHITE}), (' (Spearman 0.985)', {'color': RGBColor(0xC7, 0xD2, 0xDE)})],
    [('Learning Curve', {'color': WHITE})],
    [('Error Analysis แยกตามกลุ่ม', {'color': WHITE})],
    [('ช่วงค่า Hyperparameter ที่ค้นหา', {'color': WHITE})],
    [('Correlation Heatmap', {'color': WHITE})],
], size=12, gap=9)
notes(s, 'เปิดสไลด์สำรองจากไฟล์แยกในโฟลเดอร์ outputs/figures')

prs.save(OUT)
print(f'saved: {OUT}  ({OUT.stat().st_size/1e6:.2f} MB, {len(prs.slides.__iter__.__self__._sldIdLst)} slides)')
